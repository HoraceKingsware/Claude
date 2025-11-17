"""
Document parsing service for multiple file formats
"""
import os
from typing import Optional
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from bs4 import BeautifulSoup
import chardet


class DocumentParser:
    """Parse documents of various formats"""

    @staticmethod
    def detect_encoding(file_path: str) -> str:
        """Detect file encoding"""
        with open(file_path, 'rb') as f:
            result = chardet.detect(f.read())
            return result['encoding'] or 'utf-8'

    @staticmethod
    def parse_pdf(file_path: str) -> str:
        """Parse PDF file using PyMuPDF"""
        text_content = []
        try:
            doc = fitz.open(file_path)
            for page in doc:
                text_content.append(page.get_text())
            doc.close()
            return "\n".join(text_content)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF: {str(e)}")

    @staticmethod
    def parse_docx(file_path: str) -> str:
        """Parse DOCX file using python-docx"""
        try:
            doc = DocxDocument(file_path)
            text_content = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_content.append(para.text)
            return "\n".join(text_content)
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX: {str(e)}")

    @staticmethod
    def parse_xlsx(file_path: str) -> str:
        """Parse XLSX file using openpyxl"""
        try:
            wb = load_workbook(file_path, data_only=True)
            text_content = []
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                text_content.append(f"=== Sheet: {sheet_name} ===")
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                    if row_text.strip():
                        text_content.append(row_text)
            return "\n".join(text_content)
        except Exception as e:
            raise ValueError(f"Failed to parse XLSX: {str(e)}")

    @staticmethod
    def parse_pptx(file_path: str) -> str:
        """Parse PPTX file using python-pptx"""
        try:
            prs = Presentation(file_path)
            text_content = []
            for i, slide in enumerate(prs.slides, 1):
                text_content.append(f"=== Slide {i} ===")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_content.append(shape.text)
            return "\n".join(text_content)
        except Exception as e:
            raise ValueError(f"Failed to parse PPTX: {str(e)}")

    @staticmethod
    def parse_html(file_path: str) -> str:
        """Parse HTML file using BeautifulSoup4"""
        try:
            encoding = DocumentParser.detect_encoding(file_path)
            with open(file_path, 'r', encoding=encoding) as f:
                soup = BeautifulSoup(f.read(), 'lxml')
                # Remove script and style elements
                for script in soup(["script", "style"]):
                    script.decompose()
                # Get text
                text = soup.get_text()
                # Clean up whitespace
                lines = (line.strip() for line in text.splitlines())
                chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
                text = '\n'.join(chunk for chunk in chunks if chunk)
                return text
        except Exception as e:
            raise ValueError(f"Failed to parse HTML: {str(e)}")

    @staticmethod
    def parse_text(file_path: str) -> str:
        """Parse plain text file"""
        try:
            encoding = DocumentParser.detect_encoding(file_path)
            with open(file_path, 'r', encoding=encoding) as f:
                return f.read()
        except Exception as e:
            raise ValueError(f"Failed to parse text file: {str(e)}")

    @classmethod
    def parse(cls, file_path: str, file_type: str) -> str:
        """
        Parse document based on file type

        Args:
            file_path: Path to the file
            file_type: File extension (e.g., 'pdf', 'docx')

        Returns:
            Extracted text content
        """
        file_type = file_type.lower().lstrip('.')

        parsers = {
            'pdf': cls.parse_pdf,
            'docx': cls.parse_docx,
            'xlsx': cls.parse_xlsx,
            'xls': cls.parse_xlsx,
            'pptx': cls.parse_pptx,
            'html': cls.parse_html,
            'htm': cls.parse_html,
            'txt': cls.parse_text,
            'md': cls.parse_text,
        }

        parser = parsers.get(file_type)
        if not parser:
            raise ValueError(f"Unsupported file type: {file_type}")

        return parser(file_path)

    @staticmethod
    def get_preview(content: str, max_length: int = 500) -> str:
        """Get preview of content"""
        if len(content) <= max_length:
            return content
        return content[:max_length] + "..."
